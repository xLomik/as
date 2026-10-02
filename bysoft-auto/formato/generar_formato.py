# Genera FORMATO_CANTIDADES.xlsx: plantilla que llena el usuario para cada pedido.
# Uso: python generar_formato.py
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = PatternFill("solid", fgColor="1F4E78")
GRIS = PatternFill("solid", fgColor="F2F2F2")
BLANCO = Font(bold=True, color="FFFFFF")
FINO = Side(style="thin", color="BFBFBF")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)

wb = Workbook()

# --- Hoja 1: Cantidades (la que se llena) ---------------------------------
ws = wb.active
ws.title = "Cantidades"
cab = ["Referencia", "Cantidad", "Observacion"]
ws.append(cab)
for i, c in enumerate(ws[1], start=1):
    c.fill, c.font, c.border = AZUL, BLANCO, BORDE
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 22
ws.column_dimensions["A"].width = 45
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 40
ws.freeze_panes = "A2"
FILAS = 500
for r in range(2, FILAS + 2):
    ws.cell(r, 1).number_format = "@"          # referencia siempre texto
    ws.cell(r, 3).number_format = "@"
    for col in (1, 2, 3):
        ws.cell(r, col).border = BORDE

dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1",
                    allow_blank=True, showErrorMessage=True,
                    errorTitle="Cantidad no valida",
                    error="La cantidad debe ser un numero entero mayor o igual a 1.")
dv.add(f"B2:B{FILAS + 1}")
ws.add_data_validation(dv)
dvref = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="100",
                       allow_blank=True, showErrorMessage=True,
                       errorTitle="Referencia demasiado larga",
                       error="BySoft admite nombres de pieza de hasta 100 caracteres.")
dvref.add(f"A2:A{FILAS + 1}")
ws.add_data_validation(dvref)

# --- Hoja 2: Instrucciones -------------------------------------------------
wi = wb.create_sheet("Instrucciones")
wi.column_dimensions["A"].width = 120
lineas = [
    ("FORMATO DE CANTIDADES - Automatizacion BySoft", "titulo"),
    ("", None),
    ("1. COMO LLENAR LA HOJA 'Cantidades'", "sub"),
    ("   - Referencia: nombre EXACTO del archivo DXF, SIN la extension .dxf (ej. ESQUINERO_PLATAFORMA_IZQ_R0).", None),
    ("   - Cantidad: numero entero mayor o igual a 1.", None),
    ("   - Observacion: opcional. Se copia a 'User info 1' de la pieza en el nesteo.", None),
    ("   - Una fila por referencia. Si una referencia aparece dos veces, la herramienta SUMA las cantidades y avisa.", None),
    ("   - No dejar filas vacias en medio. No cambiar los titulos de la fila 1.", None),
    ("", None),
    ("2. COMO ORGANIZAR LA CARPETA DE DXF DEL PEDIDO", "sub"),
    ("   Una subcarpeta por cada material y espesor. Nombre de la subcarpeta:  <MATERIAL> Esp=<espesor>mm", None),
    ("      PEDIDO_XXX\\", None),
    ("         SAEJ 050 Esp=3.42mm\\      -> pieza1.dxf, pieza2.dxf ...", None),
    ("         HARDOX 450 Esp=4.5mm\\     -> pieza3.dxf ...", None),
    ("         ASTM A36 Esp=6.35mm\\      -> pieza4.dxf ...", None),
    ("   - El espesor con PUNTO decimal (3.42, no 3,42).", None),
    ("   - El MATERIAL debe estar en la tabla de equivalencias de la herramienta (se arma con el catalogo de BySoft).", None),
    ("   - Cada subcarpeta = un programa de nesteo (BySoft nestea un solo material y espesor por job).", None),
    ("", None),
    ("3. QUE HACE LA HERRAMIENTA CON ESTE ARCHIVO", "sub"),
    ("   a) Busca cada Referencia en las subcarpetas y avisa si falta algun DXF o si sobra alguno sin cantidad.", None),
    ("   b) Revisa que el nombre no exista ya en otra carpeta de BySoft (evita el error de 'piezas repetidas').", None),
    ("   c) Crea la carpeta del programa en Parts-FANALCA y PartJobs-FANALCA (con la funcion de BySoft).", None),
    ("   d) Importa los DXF con el Part Importer usando el material y espesor de cada subcarpeta.", None),
    ("   e) Genera un Excel listo para 'Importar piezas desde archivo' del Part Nester por cada subcarpeta.", None),
    ("", None),
    ("4. EJEMPLO", "sub"),
]
for texto, estilo in lineas:
    wi.append([texto])
    c = wi.cell(wi.max_row, 1)
    if estilo == "titulo":
        c.font = Font(bold=True, size=14, color="1F4E78")
    elif estilo == "sub":
        c.font = Font(bold=True, color="1F4E78")
ej0 = wi.max_row + 1
ejemplo = [("Referencia", "Cantidad", "Observacion"),
           ("ESQUINERO_PLATAFORMA_IZQ_R0", 14, "Pedido OFPT-067"),
           ("ESQUINERO_PLATAFORMA_DER_R0", 14, ""),
           ("PORTA_PATINES_IZQ_R0", 8, "")]
wi.column_dimensions["B"].width = 12
wi.column_dimensions["C"].width = 25
for i, fila in enumerate(ejemplo):
    for j, v in enumerate(fila, start=1):
        c = wi.cell(ej0 + i, j, v)
        c.border = BORDE
        if i == 0:
            c.fill, c.font = AZUL, BLANCO
        else:
            c.fill = GRIS
wb.active = 0
wb.save("FORMATO_CANTIDADES.xlsx")
print("OK FORMATO_CANTIDADES.xlsx")
