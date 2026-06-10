# -*- coding: utf-8 -*-
"""
Escribe un .xlsx con el MISMO layout posicionado del reporte Bystronic
"Lista de órdenes", a partir del diccionario que produce parser.parse_pdf().

Las posiciones (fila, columna) se replican exactas del ejemplo entregado.
Las tablas (Chapas, Piezas) crecen hacia abajo desplazando lo que sigue.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Estilos
F_TITULO  = Font(bold=True, size=14)
F_ORDEN   = Font(bold=True, size=11)
F_SECCION = Font(bold=True, size=11, color="FFFFFF")
FILL_SEC  = PatternFill("solid", fgColor="404040")
F_LABEL   = Font(bold=True)
F_HEADER  = Font(bold=True)
FILL_HDR  = PatternFill("solid", fgColor="D9D9D9")
ALIGN_L   = Alignment(horizontal="left", vertical="center")

NCOLS = 58  # A..BF


def _set(ws, r0, c0, value, font=None, fill=None, numfmt=None, off=0):
    """Coloca `value` en (r0,c0) 0-based (+off filas). 1-based para openpyxl."""
    if value is None:
        return
    cell = ws.cell(row=r0 + 1 + off, column=c0 + 1, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if numfmt:
        cell.number_format = numfmt
    cell.alignment = ALIGN_L
    return cell


def write_xlsx(d, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"

    # Columnas angostas uniformes (emula el reporte posicionado; el texto
    # desborda sobre celdas vacías a la derecha, como en el original).
    for c in range(1, NCOLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = 3.0

    g = d.get
    pieza = d.get('pieza', {}) or {}
    plano = d.get('plano', {}) or {}
    chapas = d.get('chapas', []) or []
    partes = d.get('partes', []) or []

    PCT = '0.00%'

    # ---------- Cabecera ----------
    _set(ws, 0, 0, g('titulo', 'Lista de órdenes'), F_TITULO)
    _set(ws, 1, 0, g('orden', ''), F_ORDEN)

    _set(ws, 3, 3, 'JOB', F_LABEL); _set(ws, 3, 10, g('job'))
    _set(ws, 3, 24, 'Máquina', F_LABEL); _set(ws, 3, 30, g('maquina'))

    _set(ws, 5, 3, 'Descripción', F_LABEL); _set(ws, 5, 10, g('descripcion'))
    _set(ws, 5, 24, 'Material', F_LABEL); _set(ws, 5, 30, g('material'))

    _set(ws, 7, 24, 'Grosor', F_LABEL); _set(ws, 7, 30, g('grosor')); _set(ws, 7, 35, 'mm')

    _set(ws, 9, 3, 'Info1', F_LABEL)
    _set(ws, 9, 24, 'Cant. de piezas diferentes', F_LABEL); _set(ws, 9, 37, g('cant_piezas_diferentes'))
    _set(ws, 10, 24, 'Tecnologías utilizadas', F_LABEL)
    _set(ws, 11, 24, g('tecnologias'))

    _set(ws, 13, 3, 'Info2', F_LABEL)
    _set(ws, 14, 24, 'Gas de corte', F_LABEL); _set(ws, 14, 33, g('gas_corte'))

    _set(ws, 16, 24, 'Desperdicio', F_LABEL)
    if g('desperdicio_pct') is not None:
        _set(ws, 16, 33, g('desperdicio_pct') / 100.0, numfmt=PCT)
    _set(ws, 16, 39, 'Objetivo de desperdicio ', F_LABEL)
    _set(ws, 17, 24, 'Tiempo de corte', F_LABEL); _set(ws, 17, 33, g('tiempo_corte_job'))
    _set(ws, 17, 39, 'Costes de desperdicio', F_LABEL)

    _set(ws, 18, 3, 'Info3', F_LABEL)
    _set(ws, 19, 24, 'Tiempos improd.', F_LABEL); _set(ws, 19, 33, g('tiempos_improd_job'))
    _set(ws, 19, 39, 'Peso total de desperdicio', F_LABEL); _set(ws, 19, 52, g('peso_total_desperdicio')); _set(ws, 19, 57, 'kg')
    _set(ws, 21, 24, 'Tiempo total', F_LABEL); _set(ws, 21, 33, g('tiempo_total_job'))

    # ---------- Chapas ----------
    _set(ws, 23, 0, 'Chapas', F_SECCION, FILL_SEC)
    _set(ws, 25, 0, 'Dimensión X', F_HEADER, FILL_HDR)
    _set(ws, 25, 12, 'Dimensión Y', F_HEADER, FILL_HDR)
    _set(ws, 25, 23, 'Peso', F_HEADER, FILL_HDR)
    _set(ws, 25, 33, 'Cantidad', F_HEADER, FILL_HDR)
    _set(ws, 25, 40, 'Información de artículo', F_HEADER, FILL_HDR)
    for i, ch in enumerate(chapas):
        r = 27 + i * 2
        _set(ws, r, 0, ch.get('x')); _set(ws, r, 7, 'mm')
        _set(ws, r, 12, ch.get('y')); _set(ws, r, 16, 'mm')
        _set(ws, r, 23, ch.get('peso')); _set(ws, r, 28, 'kg')
        _set(ws, r, 33, ch.get('cantidad'))
        if ch.get('info'):
            _set(ws, r, 40, ch.get('info'))
    off_a = max(0, len(chapas) - 1) * 2  # desplazamiento por chapas extra

    # ---------- Piezas (offset off_a) ----------
    _set(ws, 29, 0, 'Piezas', F_SECCION, FILL_SEC, off=off_a)
    _set(ws, 31, 15, 'Número de piezas', F_LABEL, off=off_a); _set(ws, 32, 26, pieza.get('numero_de_piezas'), off=off_a)
    _set(ws, 34, 15, 'Nombre', F_LABEL, off=off_a); _set(ws, 34, 25, pieza.get('nombre'), off=off_a)
    _set(ws, 34, 39, 'Nominal / Efectivo', F_LABEL, off=off_a); _set(ws, 34, 48, pieza.get('nominal_efectivo'), off=off_a)
    _set(ws, 36, 15, 'Descripción', F_LABEL, off=off_a); _set(ws, 36, 25, pieza.get('descripcion'), off=off_a)
    _set(ws, 36, 39, 'Dimensión X', F_LABEL, off=off_a); _set(ws, 36, 45, pieza.get('dim_x'), off=off_a); _set(ws, 36, 55, 'mm', off=off_a)
    _set(ws, 38, 15, 'Info1', F_LABEL, off=off_a)
    _set(ws, 38, 39, 'Dimensión Y', F_LABEL, off=off_a); _set(ws, 38, 45, pieza.get('dim_y'), off=off_a); _set(ws, 38, 55, 'mm', off=off_a)
    _set(ws, 40, 15, 'Info2', F_LABEL, off=off_a)
    _set(ws, 40, 39, 'Superficie efectiva', F_LABEL, off=off_a); _set(ws, 40, 45, pieza.get('superficie'), off=off_a); _set(ws, 40, 55, 'm²', off=off_a)
    _set(ws, 42, 15, 'Info3', F_LABEL, off=off_a)
    _set(ws, 42, 39, 'Peso', F_LABEL, off=off_a); _set(ws, 42, 45, pieza.get('peso'), off=off_a); _set(ws, 42, 55, 'kg', off=off_a)
    _set(ws, 44, 39, 'Perimetro', F_LABEL, off=off_a); _set(ws, 44, 45, pieza.get('perimetro'), off=off_a); _set(ws, 44, 51, 'mm', off=off_a)
    _set(ws, 46, 15, 'Info de órdenes', F_LABEL, off=off_a)
    _set(ws, 49, 15, 'Tiempo de corte', F_LABEL, off=off_a); _set(ws, 49, 24, pieza.get('tiempo_corte'), off=off_a)
    _set(ws, 49, 29, 'Tiempos improductivos', F_LABEL, off=off_a); _set(ws, 49, 40, pieza.get('tiempos_improd'), off=off_a)
    _set(ws, 49, 44, 'Tiempo total', F_LABEL, off=off_a); _set(ws, 49, 52, pieza.get('tiempo_total'), off=off_a)

    # ---------- Planos de corte (offset off_a) ----------
    _set(ws, 52, 0, 'Planos de corte', F_SECCION, FILL_SEC, off=off_a)
    _set(ws, 54, 3, 'Nombre de plano de corte', F_LABEL, off=off_a); _set(ws, 54, 14, plano.get('nombre'), off=off_a)
    _set(ws, 54, 32, 'Dimensión X del plano', F_LABEL, off=off_a); _set(ws, 54, 44, plano.get('dim_x_plano'), off=off_a); _set(ws, 54, 53, 'mm', off=off_a)
    _set(ws, 56, 3, 'Pasadas', F_LABEL, off=off_a); _set(ws, 56, 14, plano.get('pasadas'), off=off_a)
    _set(ws, 56, 32, 'Dimensión Y del plano', F_LABEL, off=off_a); _set(ws, 56, 44, plano.get('dim_y_plano'), off=off_a); _set(ws, 56, 53, 'mm', off=off_a)
    _set(ws, 58, 3, 'Desperdicio', F_LABEL, off=off_a)
    if plano.get('desperdicio_pct') is not None:
        _set(ws, 58, 14, plano.get('desperdicio_pct') / 100.0, numfmt=PCT, off=off_a)
    _set(ws, 58, 32, 'Dimensión X de chapa', F_LABEL, off=off_a); _set(ws, 58, 44, plano.get('dim_x_chapa'), off=off_a); _set(ws, 58, 53, 'mm', off=off_a)
    _set(ws, 60, 3, 'Número de piezas', F_LABEL, off=off_a); _set(ws, 60, 14, plano.get('num_piezas'), off=off_a)
    _set(ws, 60, 32, 'Dimensión Y de chapa', F_LABEL, off=off_a); _set(ws, 60, 44, plano.get('dim_y_chapa'), off=off_a); _set(ws, 60, 53, 'mm', off=off_a)
    _set(ws, 62, 3, 'Tiempo de corte', F_LABEL, off=off_a); _set(ws, 62, 10, plano.get('tiempo_corte'), off=off_a)
    _set(ws, 62, 21, 'Tiempos improductivos', F_LABEL, off=off_a); _set(ws, 62, 32, plano.get('tiempos_improd'), off=off_a)
    _set(ws, 62, 42, 'Tiempo total', F_LABEL, off=off_a); _set(ws, 62, 46, plano.get('tiempo_total'), off=off_a)

    # ---------- Tabla de piezas (offset off_a) ----------
    _set(ws, 66, 0, 'N°.', F_HEADER, FILL_HDR, off=off_a)
    _set(ws, 66, 5, 'Nombre de pieza', F_HEADER, FILL_HDR, off=off_a)
    _set(ws, 66, 16, 'Dimensión X', F_HEADER, FILL_HDR, off=off_a)
    _set(ws, 66, 26, 'Dimensión Y', F_HEADER, FILL_HDR, off=off_a)
    _set(ws, 66, 36, 'Peso', F_HEADER, FILL_HDR, off=off_a)
    _set(ws, 66, 43, 'Tiempo de corte', F_HEADER, FILL_HDR, off=off_a)
    _set(ws, 66, 53, 'Cantidad', F_HEADER, FILL_HDR, off=off_a)
    for i, pt in enumerate(partes):
        r = 68 + i * 2
        _set(ws, r, 0, pt.get('n'), off=off_a)
        _set(ws, r, 5, pt.get('nombre'), off=off_a)
        _set(ws, r, 16, pt.get('dim_x'), off=off_a); _set(ws, r, 23, 'mm', off=off_a)
        _set(ws, r, 26, pt.get('dim_y'), off=off_a); _set(ws, r, 32, 'mm', off=off_a)
        _set(ws, r, 36, pt.get('peso'), off=off_a); _set(ws, r, 42, 'kg', off=off_a)
        _set(ws, r, 43, pt.get('tiempo_corte'), off=off_a)
        _set(ws, r, 53, pt.get('cantidad'), off=off_a)
    off_b = max(0, len(partes) - 1) * 2

    # ---------- Pie (offset off_a + off_b) ----------
    off_pie = off_a + off_b
    _set(ws, 70, 1, (g('empresa', '') + '  ') if g('empresa') else None, off=off_pie)
    _set(ws, 72, 1, g('usuario'), off=off_pie)
    _set(ws, 72, 20, g('fecha'), off=off_pie)
    _set(ws, 72, 38, 'Seite 1', off=off_pie)

    wb.save(out_path)
    return out_path
